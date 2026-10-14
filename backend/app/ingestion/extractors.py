"""Source extraction: turn an uploaded file into text content units.

Supported today: PDFs (text layer via pypdf), PowerPoint decks (text + speaker
notes via python-pptx), images (local OCR via Tesseract) and plain text /
Markdown. Video is an intentional future addition.
"""

import shutil
from dataclasses import dataclass, field
from pathlib import Path

from app.core.config import get_settings
from app.core.logging import get_logger
from app.ingestion.chunking import chunk_text

logger = get_logger(__name__)
settings = get_settings()

IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp"}
TEXT_EXTENSIONS = {"txt", "md", "markdown"}
PDF_EXTENSIONS = {"pdf"}
PPTX_EXTENSIONS = {"pptx"}

_WINDOWS_TESSERACT = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")


class IngestionError(Exception):
    """Raised when a source cannot be extracted (clear, user-safe message)."""


@dataclass
class ExtractedUnit:
    content: str
    title: str | None = None
    page_number: int | None = None
    slide_number: int | None = None
    metadata: dict = field(default_factory=dict)


def detect_material_type(filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in IMAGE_EXTENSIONS:
        return "image"
    if ext in TEXT_EXTENSIONS:
        return "text"
    if ext in PDF_EXTENSIONS:
        return "pdf"
    if ext in PPTX_EXTENSIONS:
        return "pptx"
    raise IngestionError(f"Unsupported file type: .{ext}")


def _configure_tesseract() -> str:
    import pytesseract

    if settings.tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd
        return settings.tesseract_cmd
    found = shutil.which("tesseract")
    if found:
        return found
    if _WINDOWS_TESSERACT.exists():
        pytesseract.pytesseract.tesseract_cmd = str(_WINDOWS_TESSERACT)
        return str(_WINDOWS_TESSERACT)
    raise IngestionError(
        "OCR engine (Tesseract) is not installed, so image text cannot be read."
    )


def ocr_available() -> bool:
    try:
        _configure_tesseract()
        return True
    except IngestionError:
        return False


def _extract_image(path: Path) -> list[ExtractedUnit]:
    if not settings.ocr_enabled:
        raise IngestionError("OCR is disabled in settings.")
    _configure_tesseract()

    from PIL import Image
    import pytesseract

    try:
        with Image.open(path) as image:
            image = image.convert("RGB")
            size = image.size
            text = pytesseract.image_to_string(image, lang=settings.ocr_language)
    except IngestionError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise IngestionError(f"Could not read image: {type(exc).__name__}") from exc

    text = (text or "").strip()
    if not text:
        raise IngestionError("No readable text found in the image.")

    return [
        ExtractedUnit(
            content=chunk,
            title=path.stem,
            metadata={"engine": "tesseract", "width": size[0], "height": size[1]},
        )
        for chunk in chunk_text(text, settings.chunk_size, settings.chunk_overlap)
    ]


def _extract_pdf(path: Path) -> list[ExtractedUnit]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
        pages = list(reader.pages)
    except (PdfReadError, OSError, ValueError) as exc:
        raise IngestionError(
            f"Could not read PDF: {type(exc).__name__}"
        ) from exc

    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception as exc:  # noqa: BLE001
            raise IngestionError(
                "This PDF is password protected — remove the password and upload it again."
            ) from exc

    units: list[ExtractedUnit] = []
    for number, page in enumerate(pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as exc:  # noqa: BLE001 - skip unreadable page
            logger.warning("Skipping page %s of %s: %s", number, path.name, exc)
            continue
        if not text:
            continue
        for chunk in chunk_text(text, settings.chunk_size, settings.chunk_overlap):
            units.append(
                ExtractedUnit(content=chunk, title=path.stem, page_number=number)
            )

    if not units:
        raise IngestionError(
            "No readable text found in this PDF — it may be a scanned image. "
            "Upload it as an image to run OCR."
        )
    return units


def _shape_text(shape) -> list[str]:
    """Readable text from a slide shape (text frame, table or group)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
        texts: list[str] = []
        for child in shape.shapes:
            texts.extend(_shape_text(child))
        return texts

    texts = []
    if getattr(shape, "has_text_frame", False):
        text = (shape.text_frame.text or "").strip()
        if text:
            texts.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                texts.append(row_text)
    return texts


def _extract_pptx(path: Path) -> list[ExtractedUnit]:
    """Extract slide text and speaker notes from a PowerPoint (.pptx) deck."""
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise IngestionError(
            "PowerPoint reading is unavailable (python-pptx is not installed)."
        ) from exc

    try:
        presentation = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        raise IngestionError(
            f"Could not read PowerPoint file: {type(exc).__name__}"
        ) from exc

    units: list[ExtractedUnit] = []
    for number, slide in enumerate(presentation.slides, start=1):
        title = ""
        if slide.shapes.title is not None:
            title = (slide.shapes.title.text or "").strip()

        blocks: list[str] = []
        for shape in slide.shapes:
            for text in _shape_text(shape):
                if text not in blocks:
                    blocks.append(text)

        parts = list(blocks)
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                parts.append(f"Speaker notes:\n{notes}")

        slide_text = "\n".join(part for part in parts if part).strip()
        if not slide_text:
            continue

        for chunk in chunk_text(slide_text, settings.chunk_size, settings.chunk_overlap):
            units.append(
                ExtractedUnit(
                    content=chunk,
                    title=title or path.stem,
                    slide_number=number,
                    metadata={"slide": number},
                )
            )

    if not units:
        raise IngestionError(
            "No readable text found in this .pptx — it may contain only images."
        )
    return units


def _extract_text(path: Path) -> list[ExtractedUnit]:
    raw = path.read_text(encoding="utf-8", errors="ignore").strip()
    if not raw:
        raise IngestionError("The file is empty.")
    return [
        ExtractedUnit(content=chunk, title=path.stem)
        for chunk in chunk_text(raw, settings.chunk_size, settings.chunk_overlap)
    ]


def extract_units(path: Path, material_type: str) -> list[ExtractedUnit]:
    """Extract content units from a stored file of the given material type."""
    logger.info("Extracting %s from %s", material_type, path.name)
    if material_type == "image":
        units = _extract_image(path)
    elif material_type == "text":
        units = _extract_text(path)
    elif material_type == "pdf":
        units = _extract_pdf(path)
    elif material_type == "pptx":
        units = _extract_pptx(path)
    else:
        raise IngestionError(f"Unsupported material type: {material_type}")
    if not units:
        raise IngestionError("Nothing was extracted from this file.")
    return units
