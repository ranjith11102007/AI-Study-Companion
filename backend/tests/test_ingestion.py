"""Unit tests for ingestion: chunking, extractors and config."""

import httpx
import pytest

from app.core.config import get_settings
from app.ingestion import pipeline
from app.ingestion.chunking import chunk_text, sanitize_text
from app.ingestion.extractors import (
    IngestionError,
    detect_material_type,
    extract_units,
    ocr_available,
)
from app.ingestion.pipeline import Stage, _embed_all, _embed_batch, _user_safe_error

settings = get_settings()


def test_sanitize_text_strips_nul_and_control_chars() -> None:
    raw = "Operating\x00 System\r\nConcepts\x0b \x1f\x7f"
    cleaned = sanitize_text(raw)
    assert "\x00" not in cleaned
    assert "\x0b" not in cleaned
    assert "\x1f" not in cleaned
    assert "\x7f" not in cleaned
    assert "Operating System" in cleaned
    assert "\r" not in cleaned  # CRLF normalised


def test_sanitize_text_keeps_tabs_and_newlines() -> None:
    assert sanitize_text("a\tb\nc") == "a\tb\nc"
    assert sanitize_text("") == ""
    assert sanitize_text(None) == ""


def test_chunk_text_removes_nul_bytes() -> None:
    # A PDF text layer can embed NUL bytes; Postgres rejects them outright.
    chunks = chunk_text("a\x00b\x00c")
    assert chunks == ["abc"]


def test_chunk_text_empty_and_short() -> None:
    assert chunk_text("") == []
    assert chunk_text("   \n\n  ") == []
    assert chunk_text("short passage") == ["short passage"]


def test_chunk_text_respects_size_and_overlap() -> None:
    long_para = "word " * 500  # ~2500 chars, single paragraph
    chunks = chunk_text(long_para, size=400, overlap=80)
    assert len(chunks) > 1
    assert all(len(chunk) <= 400 for chunk in chunks)
    assert all(chunk for chunk in chunks)


def test_chunk_text_splits_on_paragraphs() -> None:
    para_a = "A" * 600
    text = f"{para_a}\n\n{'B' * 600}"
    chunks = chunk_text(text, size=700, overlap=0)
    assert len(chunks) == 2
    assert chunks[0].startswith("A")
    assert chunks[1].endswith("B")


def test_chunk_text_never_exceeds_size_for_long_paragraph() -> None:
    text = "x" * 5000
    chunks = chunk_text(text, size=500, overlap=100)
    assert all(len(c) <= 500 for c in chunks)
    assert "".join(chunks).count("x") >= len(text)


def test_detect_material_type() -> None:
    assert detect_material_type("scan.PNG") == "image"
    assert detect_material_type("lecture.md") == "text"
    assert detect_material_type("chapter.pdf") == "pdf"
    assert detect_material_type("os-unit-1.pptx") == "pptx"
    with pytest.raises(IngestionError):
        detect_material_type("notes.docx")


def test_extract_text_file(tmp_path) -> None:
    path = tmp_path / "notes.txt"
    body = "Photosynthesis converts light energy into chemical energy."
    path.write_text(body, encoding="utf-8")

    units = extract_units(path, "text")
    assert len(units) == 1
    assert units[0].content == body
    assert units[0].title == "notes"


def test_extract_markdown_file(tmp_path):
    path = tmp_path / "chapter.md"
    path.write_text("# Chapter 1\n\nCells are the unit of life.", encoding="utf-8")
    units = extract_units(path, "text")
    assert "Cells are the unit of life." in units[0].content


def test_extract_rejects_unsupported_type(tmp_path):
    path = tmp_path / "notes.docx"
    path.write_bytes(b"PK\x03\x04")
    with pytest.raises(IngestionError):
        extract_units(path, "docx")


def test_extract_pptx_text_and_speaker_notes(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])  # Title only
    slide.shapes.title.text = "Operating System Concepts"
    textbox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    textbox.text_frame.text = "A process is a program in execution."
    slide.notes_slide.notes_text_frame.text = "Mention the scheduler next."

    path = tmp_path / "os-unit-1.pptx"
    presentation.save(str(path))

    units = extract_units(path, "pptx")
    combined = "\n".join(unit.content for unit in units)
    assert "Operating System Concepts" in combined
    assert "program in execution" in combined
    assert "scheduler" in combined
    assert {unit.slide_number for unit in units} == {1}


def _build_pdf(page_texts: list[str]) -> bytes:
    """Build a minimal, valid single-font PDF with a text layer."""
    n = len(page_texts)
    page_ids = [3 + i for i in range(n)]
    content_ids = [3 + n + i for i in range(n)]
    font_id = 3 + 2 * n

    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: (
            "<< /Type /Pages /Kids ["
            + " ".join(f"{i} 0 R" for i in page_ids)
            + f"] /Count {n} >>"
        ).encode(),
        font_id: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    for index, text in enumerate(page_texts):
        safe = text.replace("\\", "").replace("(", "").replace(")", "")
        stream = f"BT /F1 14 Tf 50 700 Td ({safe}) Tj ET".encode()
        objects[page_ids[index]] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_ids[index]} 0 R /Resources "
            f"<< /Font << /F1 {font_id} 0 R >> >> >>"
        ).encode()
        objects[content_ids[index]] = (
            f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"
        )

    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for obj_id in sorted(objects):
        offsets[obj_id] = len(out)
        out += f"{obj_id} 0 obj\n".encode() + objects[obj_id] + b"\nendobj\n"

    xref_pos = len(out)
    count = max(objects) + 1
    out += f"xref\n0 {count}\n".encode()
    out += b"0000000000 65535 f \n"
    for obj_id in range(1, count):
        out += f"{offsets.get(obj_id, 0):010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {count} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n"
    ).encode()
    return bytes(out)


def test_extract_pdf_text_layer(tmp_path):
    path = tmp_path / "chapter.pdf"
    path.write_bytes(
        _build_pdf(
            [
                "The cell membrane regulates what enters and leaves the cell",
                "Mitochondria convert energy into ATP for the cell to use",
            ]
        )
    )

    units = extract_units(path, "pdf")
    assert len(units) == 2
    assert {u.page_number for u in units} == {1, 2}
    assert "Mitochondria" in units[1].content


def test_extract_pdf_without_text_layer_raises(tmp_path):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    path = tmp_path / "blank.pdf"
    with open(path, "wb") as handle:
        writer.write(handle)

    with pytest.raises(IngestionError) as excinfo:
        extract_units(path, "pdf")
    assert "scanned image" in str(excinfo.value).lower()


@pytest.mark.skipif(not ocr_available(), reason="Tesseract OCR is not installed")
def test_extract_image_uses_ocr(tmp_path):
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (520, 120), "white")
    draw = ImageDraw.Draw(image)
    draw.text((20, 40), "HELLO STUDY", fill="black")
    path = tmp_path / "note.png"
    image.save(path)

    units = extract_units(path, "image")
    combined = " ".join(u.content for u in units).upper()
    assert "HELLO" in combined


def test_allowed_extensions_setting():
    allowed = settings.allowed_extension_set()
    assert {"pdf", "pptx", "png", "jpg", "txt", "md"}.issubset(allowed)
    assert settings.max_upload_size_mb > 0
    assert settings.chunk_size > settings.chunk_overlap


# --------------------------------------------------------------------------
# Embedding batching + failure diagnostics (the ">100 passages" regression)
# --------------------------------------------------------------------------


class _RecordingEmbedder:
    """Stub embedding provider that records each ``embed`` call."""

    model = "stub-embed"

    def __init__(self) -> None:
        self.batches: list[list[str]] = []

    async def embed(self, texts: list[str]) -> list[list[float]]:
        self.batches.append(list(texts))
        return [[float(text)] for text in texts]


async def test_embed_all_splits_large_documents_into_batches(monkeypatch):
    """A 250-passage document must be sent as 100 + 100 + 50, in order."""
    monkeypatch.setattr(pipeline.settings, "embedding_batch_size", 100)
    provider = _RecordingEmbedder()

    vectors = await _embed_all(provider, [str(i) for i in range(250)])

    assert [len(batch) for batch in provider.batches] == [100, 100, 50]
    assert len(vectors) == 250
    # Order is preserved across the batch boundaries.
    assert [int(vector[0]) for vector in vectors] == list(range(250))


async def test_embed_all_uses_a_single_call_for_small_documents(monkeypatch):
    monkeypatch.setattr(pipeline.settings, "embedding_batch_size", 100)
    provider = _RecordingEmbedder()

    vectors = await _embed_all(provider, ["7"])

    assert [len(batch) for batch in provider.batches] == [1]
    assert vectors == [[7.0]]


async def test_embed_all_respects_a_custom_batch_size(monkeypatch):
    monkeypatch.setattr(pipeline.settings, "embedding_batch_size", 2)
    provider = _RecordingEmbedder()

    await _embed_all(provider, ["0", "1", "2", "3", "4"])

    assert [len(batch) for batch in provider.batches] == [2, 2, 1]


def _status_error(status: int, body: str) -> httpx.HTTPStatusError:
    request = httpx.Request("POST", "https://gemini.test/v1beta/openai/embeddings")
    response = httpx.Response(status, text=body, request=request)
    return httpx.HTTPStatusError(
        f"Client error '{status}'", request=request, response=response
    )


class _FlakyEmbedder:
    """Fails the first ``failures`` calls with a 429, then succeeds."""

    model = "stub-flaky"
    vector = [0.0] * 4

    def __init__(self, failures: int) -> None:
        self.failures = failures
        self.calls = 0

    async def embed(self, texts: list[str]) -> list[list[float]]:
        self.calls += 1
        if self.calls <= self.failures:
            raise _status_error(429, "quota exceeded")
        return [list(self.vector) for _ in texts]


async def test_embed_batch_retries_a_rate_limited_call(monkeypatch):
    monkeypatch.setattr(pipeline.settings, "embedding_max_attempts", 4)
    monkeypatch.setattr(pipeline.settings, "embedding_retry_base_seconds", 0.0)
    provider = _FlakyEmbedder(failures=1)

    vectors = await _embed_batch(provider, ["a"])

    assert provider.calls == 2
    assert vectors == [[0.0] * 4]


async def test_embed_batch_gives_up_after_max_attempts(monkeypatch):
    monkeypatch.setattr(pipeline.settings, "embedding_max_attempts", 3)
    monkeypatch.setattr(pipeline.settings, "embedding_retry_base_seconds", 0.0)
    provider = _FlakyEmbedder(failures=99)

    with pytest.raises(httpx.HTTPStatusError):
        await _embed_batch(provider, ["a"])

    assert provider.calls == 3


async def test_embed_batch_does_not_retry_a_rejected_request(monkeypatch):
    monkeypatch.setattr(pipeline.settings, "embedding_max_attempts", 4)
    monkeypatch.setattr(pipeline.settings, "embedding_retry_base_seconds", 0.0)

    class _Rejecting:
        async def embed(self, texts):
            self.calls += 1
            raise _status_error(400, "invalid")

        calls = 0

    provider = _Rejecting()
    with pytest.raises(httpx.HTTPStatusError):
        await _embed_batch(provider, ["a"])

    assert provider.calls == 1


async def test_embed_all_retries_rate_limits_between_batches(monkeypatch):
    monkeypatch.setattr(pipeline.settings, "embedding_batch_size", 2)
    monkeypatch.setattr(pipeline.settings, "embedding_max_attempts", 4)
    monkeypatch.setattr(pipeline.settings, "embedding_retry_base_seconds", 0.0)
    provider = _FlakyEmbedder(failures=1)

    vectors = await _embed_all(provider, [str(i) for i in range(5)])

    assert len(vectors) == 5
    assert provider.calls == 4  # 3 batches, first one retried once


def test_user_safe_error_passes_ingestion_message_through():
    message = _user_safe_error(
        IngestionError("Unsupported file type: .docx"), Stage.TEXT_EXTRACTION
    )
    assert message == "Unsupported file type: .docx"


def test_user_safe_error_explains_rate_limiting():
    message = _user_safe_error(
        _status_error(429, "quota exceeded"), Stage.EMBEDDING_GENERATION
    )
    assert "429" in message
    assert "retry" in message.lower() or "wait" in message.lower()


def test_user_safe_error_explains_a_rejected_api_key():
    message = _user_safe_error(
        _status_error(401, "unauthenticated"), Stage.EMBEDDING_GENERATION
    )
    assert "401" in message
    assert "OPENAI_API_KEY" in message


def test_user_safe_error_explains_the_batch_ceiling():
    message = _user_safe_error(
        _status_error(400, "at most 100 requests can be in one batch"),
        Stage.EMBEDDING_GENERATION,
    )
    assert "100" in message
    assert "EMBEDDING_BATCH_SIZE" in message


def test_user_safe_error_names_the_failed_stage():
    message = _user_safe_error(ValueError("boom"), Stage.VECTOR_STORAGE)
    assert "VECTOR_STORAGE" in message
    assert "Traceback" not in message


def test_user_safe_error_guides_ollama_connection_failures():
    message = _user_safe_error(
        httpx.ConnectError("All connection attempts failed"), Stage.EMBEDDING_GENERATION
    )
    assert "Ollama" in message


def test_error_detail_includes_the_provider_body():
    detail = pipeline._error_detail(
        _status_error(400, "requests: at most 100 requests can be in one batch")
    )
    assert "400" in detail
    assert "at most 100" in detail

