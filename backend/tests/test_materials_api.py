"""Integration tests for material upload, indexing and per-user isolation."""

import uuid

import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy import delete

from app.db.session import SessionLocal
from app.ingestion import pipeline
from app.main import app
from app.models.user import User

PASSWORD = "secret123"


class _FakeEmbeddingProvider:
    """Deterministic embeddings so tests never depend on a running Ollama."""

    model = "test-embed"

    async def embed(self, texts: list[str]) -> list[list[float]]:
        return [[0.0] * 768 for _ in texts]


@pytest.fixture
def fake_embeddings(monkeypatch):
    monkeypatch.setattr(pipeline, "get_embedding_provider", _FakeEmbeddingProvider)


async def _make_user(client: AsyncClient, email: str) -> dict:
    await client.post(
        "/api/auth/signup",
        json={"name": "Test", "email": email, "password": PASSWORD},
    )
    token = (
        await client.post(
            "/api/auth/signin", json={"email": email, "password": PASSWORD}
        )
    ).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


async def _cleanup(email: str) -> None:
    async with SessionLocal() as db:
        await db.execute(delete(User).filter(User.email == email))
        await db.commit()


@pytest.mark.asyncio
async def test_upload_text_material_is_extracted_and_indexed(fake_embeddings):
    email = f"mat-{uuid.uuid4().hex[:8]}@example.com"
    transport = ASGITransport(app=app)
    try:
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            headers = await _make_user(client, email)
            course = (
                await client.post(
                    "/api/courses/", json={"title": "Biology"}, headers=headers
                )
            ).json()

            response = await client.post(
                f"/api/courses/{course['id']}/materials",
                files={
                    "file": (
                        "notes.txt",
                        b"Photosynthesis converts light into chemical energy.",
                        "text/plain",
                    )
                },
                headers=headers,
            )
            assert response.status_code == 201, response.text
            material = response.json()
            assert material["material_type"] == "text"
            assert material["original_filename"] == "notes.txt"
            assert material["processing_status"] == "pending"

            listing = (
                await client.get(
                    f"/api/courses/{course['id']}/materials", headers=headers
                )
            ).json()
            assert len(listing) == 1

            detail = (
                await client.get(f"/api/materials/{material['id']}", headers=headers)
            ).json()
            assert detail["processing_status"] == "ready"
            assert detail["content_unit_count"] == 1
    finally:
        await _cleanup(email)


@pytest.mark.asyncio
async def test_upload_rejects_unsupported_and_empty_files():
    email = f"mat-{uuid.uuid4().hex[:8]}@example.com"
    transport = ASGITransport(app=app)
    try:
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            headers = await _make_user(client, email)
            course = (
                await client.post(
                    "/api/courses/", json={"title": "Chem"}, headers=headers
                )
            ).json()
            url = f"/api/courses/{course['id']}/materials"

            bad_ext = await client.post(
                url,
                files={"file": ("deck.docx", b"PK\x03\x04", "application/octet-stream")},
                headers=headers,
            )
            assert bad_ext.status_code == 400

            empty = await client.post(
                url,
                files={"file": ("empty.txt", b"", "text/plain")},
                headers=headers,
            )
            assert empty.status_code == 400
    finally:
        await _cleanup(email)


@pytest.mark.asyncio
async def test_upload_pptx_material_is_extracted_and_indexed(fake_embeddings):
    from io import BytesIO

    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Cell Biology"
    textbox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    textbox.text_frame.text = "Mitochondria produce ATP for the cell."
    buffer = BytesIO()
    presentation.save(buffer)

    email = f"mat-{uuid.uuid4().hex[:8]}@example.com"
    transport = ASGITransport(app=app)
    try:
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            headers = await _make_user(client, email)
            course = (
                await client.post(
                    "/api/courses/", json={"title": "Bio slides"}, headers=headers
                )
            ).json()

            response = await client.post(
                f"/api/courses/{course['id']}/materials",
                files={
                    "file": (
                        "cells.pptx",
                        buffer.getvalue(),
                        "application/vnd.openxmlformats-officedocument."
                        "presentationml.presentation",
                    )
                },
                headers=headers,
            )
            assert response.status_code == 201, response.text
            material = response.json()
            assert material["material_type"] == "pptx"

            detail = (
                await client.get(f"/api/materials/{material['id']}", headers=headers)
            ).json()
            assert detail["processing_status"] == "ready", detail
            assert detail["content_unit_count"] >= 1
    finally:
        await _cleanup(email)


@pytest.mark.asyncio
async def test_upload_enforces_size_limit(monkeypatch):
    from app.services import materials as material_service

    monkeypatch.setattr(material_service.settings, "max_upload_size_mb", 0)

    email = f"mat-{uuid.uuid4().hex[:8]}@example.com"
    transport = ASGITransport(app=app)
    try:
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            headers = await _make_user(client, email)
            course = (
                await client.post(
                    "/api/courses/", json={"title": "Phys"}, headers=headers
                )
            ).json()
            response = await client.post(
                f"/api/courses/{course['id']}/materials",
                files={"file": ("big.txt", b"x" * 64, "text/plain")},
                headers=headers,
            )
            assert response.status_code == 413
    finally:
        await _cleanup(email)


@pytest.mark.asyncio
async def test_material_access_is_owner_scoped(fake_embeddings):
    email_a = f"mat-{uuid.uuid4().hex[:8]}@example.com"
    email_b = f"mat-{uuid.uuid4().hex[:8]}@example.com"
    transport = ASGITransport(app=app)
    try:
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            headers_a = await _make_user(client, email_a)
            headers_b = await _make_user(client, email_b)

            course = (
                await client.post(
                    "/api/courses/", json={"title": "Private"}, headers=headers_a
                )
            ).json()
            material = (
                await client.post(
                    f"/api/courses/{course['id']}/materials",
                    files={"file": ("secret.txt", b"private notes", "text/plain")},
                    headers=headers_a,
                )
            ).json()

            other_material = await client.get(
                f"/api/materials/{material['id']}", headers=headers_b
            )
            assert other_material.status_code == 404

            other_listing = await client.get(
                f"/api/courses/{course['id']}/materials", headers=headers_b
            )
            assert other_listing.status_code == 404

            other_upload = await client.post(
                f"/api/courses/{course['id']}/materials",
                files={"file": ("x.txt", b"nope", "text/plain")},
                headers=headers_b,
            )
            assert other_upload.status_code == 404
    finally:
        await _cleanup(email_a)
        await _cleanup(email_b)


@pytest.mark.asyncio
async def test_upload_text_with_nul_bytes_is_indexed(fake_embeddings):
    """Regression: NUL bytes in extracted text used to fail the whole upload."""
    email = f"mat-{uuid.uuid4().hex[:8]}@example.com"
    transport = ASGITransport(app=app)
    try:
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            headers = await _make_user(client, email)
            course = (
                await client.post(
                    "/api/courses/", json={"title": "OS Notes"}, headers=headers
                )
            ).json()
            response = await client.post(
                f"/api/courses/{course['id']}/materials",
                files={
                    "file": (
                        "os.txt",
                        b"Operating\x00 System Concepts\x00",
                        "text/plain",
                    )
                },
                headers=headers,
            )
            assert response.status_code == 201, response.text
            material = response.json()

            detail = (
                await client.get(f"/api/materials/{material['id']}", headers=headers)
            ).json()
            assert detail["processing_status"] == "ready", detail
            assert detail["content_unit_count"] == 1
            assert not detail.get("processing_error")
    finally:
        await _cleanup(email)


@pytest.mark.asyncio
async def test_reprocess_material_reindexes(fake_embeddings):
    email = f"mat-{uuid.uuid4().hex[:8]}@example.com"
    transport = ASGITransport(app=app)
    try:
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            headers = await _make_user(client, email)
            course = (
                await client.post(
                    "/api/courses/", json={"title": "Reprocess"}, headers=headers
                )
            ).json()
            material = (
                await client.post(
                    f"/api/courses/{course['id']}/materials",
                    files={
                        "file": (
                            "notes.txt",
                            b"Cells are the unit of life.",
                            "text/plain",
                        )
                    },
                    headers=headers,
                )
            ).json()

            again = await client.post(
                f"/api/materials/{material['id']}/reprocess", headers=headers
            )
            assert again.status_code == 200, again.text

            detail = (
                await client.get(f"/api/materials/{material['id']}", headers=headers)
            ).json()
            assert detail["processing_status"] == "ready", detail
            assert detail["content_unit_count"] == 1
    finally:
        await _cleanup(email)


@pytest.mark.asyncio
async def test_reprocess_is_owner_scoped(fake_embeddings):
    email_a = f"mat-{uuid.uuid4().hex[:8]}@example.com"
    email_b = f"mat-{uuid.uuid4().hex[:8]}@example.com"
    transport = ASGITransport(app=app)
    try:
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            headers_a = await _make_user(client, email_a)
            headers_b = await _make_user(client, email_b)
            course = (
                await client.post(
                    "/api/courses/", json={"title": "Mine"}, headers=headers_a
                )
            ).json()
            material = (
                await client.post(
                    f"/api/courses/{course['id']}/materials",
                    files={"file": ("a.txt", b"hello", "text/plain")},
                    headers=headers_a,
                )
            ).json()

            forbidden = await client.post(
                f"/api/materials/{material['id']}/reprocess", headers=headers_b
            )
            assert forbidden.status_code == 404
    finally:
        await _cleanup(email_a)
        await _cleanup(email_b)
